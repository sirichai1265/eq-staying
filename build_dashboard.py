# -*- coding: utf-8 -*-
"""
สร้าง EQ STAYING Dashboard (AV / DMG / LONG STAYING + ตู้ปล่อยออก Gate Out)

วิธีใช้ (ใส่ไฟล์ STAYING และ/หรือ GATE OUT กี่ไฟล์ก็ได้ — แยกประเภทจากชื่อไฟล์):
    python build_dashboard.py "C:/Users/HAL-USER/Desktop/10-15-STAYING.xls"
    python build_dashboard.py "C:/Users/HAL-USER/Desktop/10-15-GATE OUT.xls"
    python build_dashboard.py "...STAYING.xls" "...GATE OUT.xls"
    python build_dashboard.py                      # สร้างใหม่จากไฟล์ที่เก็บไว้แล้ว
ผลลัพธ์: index.html
และเก็บไฟล์ของแต่ละวันไว้ในโฟลเดอร์วันที่ เช่น 2026-10-15/
    STAYING   -> ไฟล์ต้นฉบับ + EQ_STAYING_Dashboard_<วัน>.html + EQ_STAYING_Summary_<วัน>.xlsx
    GATE OUT  -> ไฟล์ต้นฉบับ + GATE_OUT_Report_<วัน>.xlsx (ตู้ปล่อยออก OEP/OEV/OEW พร้อมวันค้างก่อนปล่อย)
"""
import base64
import glob
import json
import os
import re
import shutil
import sys
from datetime import datetime

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "index.html")

AV_CODES = {"IEC", "IED", "IEP", "IER", "IEW", "VED"}
DMG_CODES = {"OER"}
BKK_OVERRIDE_LOCATIONS = {"LCH55", "LCH54"}
REEFER_TYPES = {"22RE", "45RE"}


RELEASE_MOVES = {"OEP", "OEV", "OEW"}   # ปล่อยตู้ออกจากลาน
DAY_DIR = re.compile(r"\d{4}-\d{2}-\d{2}")


def is_gateout(path):
    return "GATE OUT" in os.path.basename(path).upper().replace("_", " ")


def xls_in(folder, gate):
    return sorted(f for f in glob.glob(os.path.join(folder, "*.xls*"))
                  if not os.path.basename(f).startswith(("~$", "EQ_STAYING_", "GATE_OUT_")) and is_gateout(f) == gate
                  and (gate or "STAYING" in os.path.basename(f).upper()))


def day_folders():
    return sorted(d for d in os.listdir(HERE) if DAY_DIR.fullmatch(d) and os.path.isdir(os.path.join(HERE, d)))




def latest_archived_staying():
    snaps = staying_snapshots()
    if snaps:
        return snaps[max(snaps)]
    for d in reversed(day_folders()):
        files = xls_in(os.path.join(HERE, d), gate=False)
        if files:
            return files[-1]
    sys.exit("ไม่พบไฟล์ STAYING ที่เก็บไว้")


def group_of(code):
    if code in AV_CODES:
        return "AV"
    if code in DMG_CODES:
        return "DMG"
    return "OTHER"


def area_of(location, area):
    if location in BKK_OVERRIDE_LOCATIONS:
        return "BKK"
    if location.startswith("BKK"):
        return "BKK"
    if location.startswith("LCH"):
        return "LCH"
    return area or "-"


def clean(v):
    if pd.isna(v):
        return ""
    return str(v).strip()


def fmt_vl(v):
    """202604211133 -> 2026-04-21 11:33"""
    if pd.isna(v):
        return ""
    t = str(int(v))
    return f"{t[:4]}-{t[4:6]}-{t[6:8]} {t[8:10]}:{t[10:12]}" if len(t) == 12 else t


def logo_data_uri(name="heunga_logo.png", mime="image/png"):
    path = os.path.join(HERE, name)
    if not os.path.exists(path):
        return ""
    with open(path, "rb") as f:
        return f"data:{mime};base64," + base64.b64encode(f.read()).decode()


def main():
    args = sys.argv[1:]
    gate_files = [a for a in args if is_gateout(a)]
    stay_files = [a for a in args if not is_gateout(a)]

    # 1) ไฟล์ GATE OUT: ย้ายเข้าโฟลเดอร์ตามวันที่ Gate แล้วทำรายงาน
    for g in gate_files:
        archive_gateout(g)

    # 2) STAYING: ใช้ไฟล์ที่ให้มา หรือไฟล์ล่าสุดที่เก็บไว้
    if stay_files:
        for s in stay_files[:-1]:
            build_staying(s, archive_it=True, write_index=False)
        build_staying(stay_files[-1], archive_it=True)
    else:
        build_staying(latest_archived_staying(), archive_it=False)


def load_staying_rows(src):
    df = pd.read_excel(src)
    # ตัดแถวสรุปท้ายไฟล์ / แถวว่าง
    df = df[df["Location"].notna() & df["Size/Type"].notna()].copy()
    # ตัดตู้ที่จัดปล่อยแล้ว (Move Code OEV / OEP / OEW) ออกจากข้อมูล STAYING
    df = df[~df["Move Code"].astype(str).str.strip().isin(RELEASE_MOVES)]

    rows = []
    for r in df.itertuples(index=False):
        row = dict(zip(df.columns, r))
        loc = clean(row["Location"])
        code = clean(row["Move Code"])
        size = clean(row["Size/Type"])
        by = row["Built Year"]
        rows.append({
            "cn": clean(row["Container No"]),
            "st": size,
            "loc": loc,
            "ar": area_of(loc, clean(row["Area"])),
            "mc": code,
            "gp": group_of(code),
            "d": int(row["Days"]) if pd.notna(row["Days"]) else 0,
            "eq": clean(row["EQ Date"])[:10],
            "by": int(by) if pd.notna(by) else None,
            "rf": clean(row["RF Brand"]).upper() if size in REEFER_TYPES else "",
            "vl": fmt_vl(row["VL Date"]),
            "fe": clean(row["Full/Empty"]),
        })

    rows.sort(key=lambda x: -x["d"])
    return rows


def snap_date(path, last_eq):
    """วันที่ของไฟล์ STAYING (ข้อมูลเช้าวันนั้น): จากโฟลเดอร์ YYYY-MM-DD / M-D หรือชื่อไฟล์ M-D-... ไม่งั้น = EQ Date ล่าสุด + 1"""
    from datetime import timedelta
    parent = os.path.basename(os.path.dirname(os.path.abspath(path)))
    if DAY_DIR.fullmatch(parent):
        return parent
    m = re.fullmatch(r"(\d{1,2})-(\d{1,2})", parent) or re.match(r"(\d{1,2})-(\d{1,2})-", os.path.basename(path))
    if m:
        return f"2026-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    return (datetime.strptime(last_eq, "%Y-%m-%d") + timedelta(days=1)).strftime("%Y-%m-%d") if last_eq else ""


def build_staying(src, archive_it=True, write_index=True):
    rows = load_staying_rows(src)
    eq_dates = [x["eq"] for x in rows if x["eq"]]
    meta = {
        "source": os.path.basename(src),
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "lastEq": max(eq_dates) if eq_dates else "",
    }
    meta["snap"] = snap_date(src, meta["lastEq"])
    meta["holidays"] = YARD_HOLIDAYS
    gate = load_gateout_history()
    daily = long_daily()

    print(f"Source : {meta['source']}  ({len(rows)} containers)")
    for g in ("AV", "DMG", "OTHER"):
        print(f"  {g:5}: {sum(1 for x in rows if x['gp'] == g)}")
    print(f"  >=100 days: {sum(1 for x in rows if x['d'] >= 100)}   >=200 days: {sum(1 for x in rows if x['d'] >= 200)}")
    print(f"Gate out history: {len(gate)} ตู้ปล่อยออก ({', '.join(sorted({g['gd'] for g in gate})) or '-'})")
    if write_index or archive_it:
        with open(os.path.join(HERE, "dashboard_template.html"), encoding="utf-8") as f:
            html = f.read()
        html = html.replace("/*__DATA__*/null", json.dumps(rows, ensure_ascii=False, separators=(",", ":")))
        html = html.replace("/*__GATE__*/null", json.dumps(gate, ensure_ascii=False, separators=(",", ":")))
        html = html.replace("/*__LONGDAILY__*/null", json.dumps(daily, ensure_ascii=False, separators=(",", ":")))
        html = html.replace("__LOGO__", logo_data_uri())
        html = html.replace("__BANNER__", logo_data_uri("heunga_banner.jpg", "image/jpeg"))
        html = html.replace("/*__META__*/null", json.dumps(meta, ensure_ascii=False))
        with open(OUT, "w", encoding="utf-8") as f:
            f.write(html)
        print(f"Output : {OUT}")
    if archive_it:
        src = archive(src, rows, meta)
    if write_index:
        write_gate_summaries(rows, meta, gate, daily)


# ===================== GATE OUT (ตู้ปล่อยออก) =====================
def read_gateout(path):
    g = pd.read_excel(path, header=0)
    g = g.rename(columns={g.columns[0]: "Loc"})
    g = g[g["Container No"].notna() & g["Move"].notna()].copy()
    return g


def gate_day(path):
    d = read_gateout(path)["GateDate"].dropna().astype(str).str[:10]
    return max(d) if len(d) else datetime.now().strftime("%Y-%m-%d")


def staying_lookup(gate_date):
    """ข้อมูล EQ Date / Move Code ก่อนปล่อย จากไฟล์ STAYING ที่เก็บไว้ (วันเดียวกันหรือก่อนหน้า ใหม่สุดก่อน)"""
    found = {}
    for d in reversed(day_folders()):
        if d > gate_date:
            continue
        for f in xls_in(os.path.join(HERE, d), gate=False):
            cols = ["Container No", "EQ Date", "Move Code", "Location"]
            s = pd.read_excel(f, usecols=cols).dropna(subset=["Container No", "EQ Date"])[cols]
            for cn, eq, mc, loc in s.itertuples(index=False):
                found.setdefault(str(cn).strip(), (str(eq)[:10], clean(mc), clean(loc)))
    return found


def gateout_rows(path):
    g = read_gateout(path)
    g = g[g["Move"].isin(RELEASE_MOVES)]
    if g.empty:
        return []
    look = staying_lookup(gate_day(path))
    out = []
    for r in g.itertuples(index=False):
        row = dict(zip(g.columns, r))
        cn = clean(row["Container No"])
        gd = clean(row["GateDate"])[:10]
        loc = clean(row["Loc"])
        eq, pmc, sloc = look.get(cn, ("", "", ""))
        days = (datetime.strptime(gd, "%Y-%m-%d") - datetime.strptime(eq, "%Y-%m-%d")).days if eq and eq <= gd else None
        t = clean(row["GateTime"]).split(".")[0].zfill(4)
        by = row.get("Built Year")
        out.append({
            "cn": cn, "mv": clean(row["Move"]), "gd": gd, "gt": f"{t[:2]}:{t[2:]}",
            "loc": loc, "ar": area_of(loc, ""), "st": clean(row["TPSZ"]),
            "eq": eq, "d": days, "pmc": pmc, "pgp": group_of(pmc) if pmc else "",
            "by": int(by) if pd.notna(by) and str(by).strip().isdigit() else None,
            "pod": clean(row.get("POD")), "voy": clean(row.get("Voyage")),
        })
    out.sort(key=lambda x: -(x["d"] if x["d"] is not None else -1))
    return out


# ===== ตู้ปล่อยออกจากรีพอตลาน (GATE MOVE DATABASE) =====
GATE_DB = os.path.join(os.path.dirname(HERE), "GATE MOVE", "GATE MOVE DATABASE.xlsx")
STOCK_DAILY = os.path.join(os.path.dirname(HERE), "STOCK DAILY")
GATE_FROM, GATE_TO = "2026-10-01", "9999-12-31"     # ช่วงวันที่เก็บข้อมูลตู้ปล่อยออก (ถึงวันล่าสุดในฐานข้อมูล)
YARD_HOLIDAYS = {"2026-10-23": "ลานหยุด (วันปิยมหาราช)"}   # วันที่ลานหยุด — ไม่มี GATE OUT
SNAP_CACHE = os.path.join(HERE, ".staying_cache.json")


def staying_snapshots():
    """ไฟล์ STAYING รายวัน: {วันที่: path} จาก STOCK DAILY\\<m-d>\\ และโฟลเดอร์วันที่ใน EQ STAYING"""
    snaps = {}
    if os.path.isdir(STOCK_DAILY):
        for d in os.listdir(STOCK_DAILY):
            m = re.fullmatch(r"(\d{1,2})-(\d{1,2})", d)
            if not m:
                continue
            files = xls_in(os.path.join(STOCK_DAILY, d), gate=False)
            if files:
                snaps[f"2026-{int(m.group(1)):02d}-{int(m.group(2)):02d}"] = files[-1]
    for d in day_folders():
        files = xls_in(os.path.join(HERE, d), gate=False)
        if files:
            snaps.setdefault(d, files[-1])
    return dict(sorted(snaps.items()))


def load_snapshot(path, cache):
    """{เบอร์ตู้: [EQ Date, Move Code, Location, Size/Type]}"""
    key = f"v2|{path}|{os.path.getmtime(path)}"
    if key not in cache:
        cols = ["Container No", "EQ Date", "Move Code", "Location", "Size/Type"]
        s = pd.read_excel(path, usecols=cols).dropna(subset=["Container No", "EQ Date"])[cols]
        cache[key] = {str(cn).strip(): [str(eq)[:10], clean(mc), clean(loc), clean(st)]
                      for cn, eq, mc, loc, st in s.itertuples(index=False)}
    return cache[key]


SNAPS = {}      # {วันที่: {เบอร์ตู้: [...]}}  โหลดใน load_gateout_history
MOVES = {}      # {เบอร์ตู้: [(วันที่, เวลา, "IN"/"OUT", ลาน, Location), ...]}


def released_between(cn, start, end):
    """ตู้ถูกปล่อยออกระหว่าง start..end (ความเคลื่อนไหวล่าสุดในช่วงนี้คือ GATE OUT) -> คืน move นั้น"""
    last = None
    for mv in MOVES.get(cn, ()):
        if start <= mv[0] <= end and (last is None or mv[:2] > last[:2]):
            last = mv
    return last if last and last[2] == "OUT" else None


def fmt_time(v):
    """'08:19:05' / '9.59' / '13.4' -> '08:19' / '09:59' / '13:40'"""
    m = re.search(r"(\d{1,2})(?:[:.](\d{1,2}))?", clean(v))
    return f"{int(m.group(1)):02d}:{(m.group(2) or '00').ljust(2, '0')}" if m else clean(v)


def load_gateout_history():
    """GATE OUT จากรีพอตลาน (GATE MOVE DATABASE) ช่วง GATE_FROM–GATE_TO + วันค้างจาก STAYING วันก่อนปล่อย"""
    if not os.path.exists(GATE_DB):
        print(f"!! ไม่พบ {GATE_DB}")
        return []
    allmv = pd.read_excel(GATE_DB, "Container Moves")
    allmv["gd"] = allmv["Date"].astype(str).str[:10]
    MOVES.clear()
    for cn, gd, gt, mv, yard, loc in allmv[["Container No", "gd", "Gate Time", "Move", "Yard", "Location"]].itertuples(index=False):
        MOVES.setdefault(clean(cn), []).append((gd, fmt_time(gt), "OUT" if mv == "GATE OUT" else "IN", clean(yard), clean(loc)))
    # GATE IN ล่าสุดของแต่ละตู้ (ไว้เช็คว่าปล่อยออกแล้วกลับเข้าลานอีกหรือไม่)
    last_in = {}
    for cn, gd, gt in allmv[allmv["Move"] == "GATE IN"][["Container No", "gd", "Gate Time"]].itertuples(index=False):
        k = (gd, fmt_time(gt))
        cn = clean(cn)
        if k > last_in.get(cn, ("", "")):
            last_in[cn] = k
    db = allmv[(allmv["Move"] == "GATE OUT") & (allmv["gd"] >= GATE_FROM) & (allmv["gd"] <= GATE_TO)].copy()
    try:
        with open(SNAP_CACHE, encoding="utf-8") as f:
            cache = json.load(f)
    except (OSError, ValueError):
        cache = {}
    snaps = staying_snapshots()
    loaded = {d: load_snapshot(p, cache) for d, p in snaps.items() if d <= GATE_TO}
    SNAPS.clear()
    SNAPS.update(loaded)
    with open(SNAP_CACHE, "w", encoding="utf-8") as f:
        json.dump({k: v for k, v in cache.items() if k.split("|")[0] in snaps.values()}, f)
    days_desc = sorted(loaded, reverse=True)

    out = []
    for r in db.itertuples(index=False):
        row = dict(zip(db.columns, r))
        cn, gd, loc = clean(row["Container No"]), row["gd"], clean(row["Location"])
        eq, pmc = "", ""
        for d in days_desc:                      # STAYING ล่าสุดที่ไม่เกินวันปล่อย ที่มีตู้นี้
            if d <= gd and cn in loaded[d]:
                eq, pmc = loaded[d][cn][:2]
                break
        days = (datetime.strptime(gd, "%Y-%m-%d") - datetime.strptime(eq, "%Y-%m-%d")).days if eq and eq <= gd else None
        gt = fmt_time(row["Gate Time"])
        out.append({
            "cn": cn, "gd": gd, "gt": gt, "loc": loc, "ar": area_of(loc, ""),
            "back": last_in.get(cn, ("", "")) > (gd, gt),      # กลับเข้าลานหลังปล่อยออก
            "yard": clean(row["Yard"]), "typ": clean(row["Type"]), "st": clean(row["TPSZ"]),
            "eq": eq, "d": days, "pmc": pmc, "pgp": group_of(pmc) if pmc else "",
        })
    out.sort(key=lambda x: (x["gd"], x["gt"]), reverse=True)
    return out


def long_daily():
    """ทุกวันที่มี GATE OUT (ตั้งแต่ GATE_FROM): LONG STAYING ณ วันนั้นจาก STAYING ล่าสุดที่ไม่เกินวันนั้น
    - ต้นวัน = ตู้ใน snapshot ที่อายุ ≥100 วัน ณ วันนั้น (ไม่รวม OEV/OEP/OEW) และยังไม่ถูกปล่อยก่อนวันนั้น
    - ปล่อยวันนี้ = GATE OUT วันนั้นตามรีพอตลาน · คงเหลือ = ต้นวัน − ปล่อยวันนี้"""
    if not SNAPS:
        return []
    from datetime import timedelta
    days = {mv[0] for mvs in MOVES.values() for mv in mvs if mv[2] == "OUT" and GATE_FROM <= mv[0] <= GATE_TO}
    days = sorted(days | {h for h in YARD_HOLIDAYS if days and min(days) <= h <= max(days)})
    out = []
    for d in days:
        sd = max((x for x in SNAPS if x <= d), default=None)
        if not sd:
            continue
        snap = SNAPS[sd]
        dt = datetime.strptime(d, "%Y-%m-%d")
        prev = (dt - timedelta(days=1)).strftime("%Y-%m-%d")
        start, gone_today, in_yard = [], [], 0
        for cn, (eq, mc, loc, st) in snap.items():
            if mc in RELEASE_MOVES or not eq:
                continue
            if released_between(cn, sd, prev):
                continue                                 # ปล่อยไปแล้วก่อนวันนี้
            in_yard += 1
            age = (dt - datetime.strptime(eq, "%Y-%m-%d")).days
            if age < 100:
                continue
            rec = {"cn": cn, "loc": loc, "st": st, "mc": mc, "eq": eq, "d": age}
            start.append(rec)
            g = released_between(cn, d, d)
            if g:
                gone_today.append(dict(rec, gt=g[1], yard=g[3]))
        locs = sorted({r["loc"] for r in start})
        out.append({
            "d": d, "snap": sd, "inYard": in_yard, "note": YARD_HOLIDAYS.get(d, ""),
            "start": len(start), "out": len(gone_today), "end": len(start) - len(gone_today),
            "byLoc": {l: [sum(r["loc"] == l for r in start), sum(r["loc"] == l for r in gone_today)] for l in locs},
            "gone": sorted(gone_today, key=lambda r: -r["d"]),
            "remain": sorted([r for r in start if r["cn"] not in {g["cn"] for g in gone_today}], key=lambda r: -r["d"]),
        })
    write_long_daily_xlsx(out)
    for x in out:      # ในหน้าเว็บเก็บแค่ตัวเลข + รายการที่ปล่อยออก (รายการคงเหลืออยู่ใน Excel)
        x.pop("remain")
    if out:
        print(f"LONG STAYING รายวัน {out[0]['d']} – {out[-1]['d']}: ล่าสุด ต้นวัน {out[-1]['start']} · ปล่อย {out[-1]['out']} · คงเหลือ {out[-1]['end']}")
    return out


def write_long_daily_xlsx(daily):
    """EQ STAYING/LONG STAYING/LONG_STAYING_Daily_<จาก>_<ถึง>.xlsx"""
    if not daily:
        return
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="BFC5CE")
    border = Border(top=thin, left=thin, bottom=thin, right=thin)

    def head(ws, row, vals, color="1F5FBF"):
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=color)
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border = border

    def put(ws, row, vals, bold=False, fill=None):
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v)
            c.border = border
            if bold:
                c.font = Font(bold=True)
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)

    locs = sorted({l for x in daily for l in x["byLoc"]})
    wb = Workbook()
    ws = wb.active
    ws.title = "สรุปรายวัน"
    ws["A1"] = f"LONG STAYING (≥100 วัน · ไม่รวม OEV/OEP/OEW) รายวัน {daily[0]['d']} – {daily[-1]['d']}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("ต้นวัน = ตู้ใน STAYING ล่าสุดที่ไม่เกินวันนั้น (STOCK DAILY) อายุ ≥100 วัน ณ วันนั้น และยังไม่ถูกปล่อย · "
                "ปล่อยวันนี้ = GATE OUT ตามรีพอตลาน (GATE MOVE) · คงเหลือ = ต้นวัน − ปล่อยวันนี้")
    ws["A2"].font = Font(color="6B7684")
    head(ws, 4, ["วันที่", "STAYING ที่ใช้", "ตู้ในลาน (ไม่รวม OE)", "LONG ต้นวัน", "ปล่อยออกวันนี้", "LONG คงเหลือ"]
         + [f"{l} คงเหลือ" for l in locs] + ["หมายเหตุ"])
    for i, x in enumerate(daily, 5):
        put(ws, i, [x["d"], x["snap"], x["inYard"], x["start"], x["out"], x["end"]]
            + [x["byLoc"].get(l, [0, 0])[0] - x["byLoc"].get(l, [0, 0])[1] for l in locs] + [x.get("note", "")],
            fill="FEF9C3" if x.get("note") else None)
        ws.cell(i, 5).font = Font(bold=True, color="0F766E")
        ws.cell(i, 6).font = Font(bold=True, color="B91C1C")
    tot = [sum(x["out"] for x in daily)]
    put(ws, len(daily) + 5, ["รวมปล่อยออก", "", "", "", tot[0], ""], bold=True, fill="EEF1F5")
    for col, w in zip("ABCDEFGHIJKLMN", [12, 13, 12, 11, 12, 12] + [12] * 8):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("ปล่อยออกแต่ละลาน")
    head(ws2, 1, ["วันที่"] + locs + ["รวม"], "0F766E")
    for i, x in enumerate(daily, 2):
        vals = [x["byLoc"].get(l, [0, 0])[1] for l in locs]
        put(ws2, i, [x["d"]] + vals + [sum(vals)])
    put(ws2, len(daily) + 2, ["รวม"] + [sum(x["byLoc"].get(l, [0, 0])[1] for x in daily) for l in locs] + [tot[0]], bold=True, fill="EEF1F5")
    ws2.column_dimensions["A"].width = 12

    def list_sheet(name, rows, color, with_gate):
        sh = wb.create_sheet(name)
        hd = ["No.", "วันที่", "Location", "Container No", "Size/Type", "วันค้าง", "EQ Date", "Move Code"] + (["GATE OUT", "ลาน"] if with_gate else [])
        head(sh, 1, hd, color)
        for i, (d, r) in enumerate(rows, 1):
            put(sh, i + 1, [i, d, r["loc"], r["cn"], r["st"], r["d"], r["eq"], r["mc"]] + ([r["gt"], r["yard"]] if with_gate else []),
                fill="FFF59D" if r["d"] >= 200 else ("F0FDFA" if with_gate else "FEE2E2"))
            sh.cell(i + 1, 6).font = Font(bold=True, color="B91C1C")
        for col, w in zip("ABCDEFGHIJ", [6, 12, 10, 16, 10, 8, 12, 11, 9, 11]):
            sh.column_dimensions[col].width = w
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = f"A1:{chr(64 + len(hd))}{len(rows) + 1}"

    list_sheet("ตู้ค้างนานที่ปล่อยออก", [(x["d"], r) for x in daily for r in x["gone"]], "0F766E", True)
    last = daily[-1]
    list_sheet(f"คงเหลือ {last['d']}", [(last["d"], r) for r in last["remain"]], "B91C1C", False)
    outdir = os.path.join(HERE, "LONG STAYING")
    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, f"LONG_STAYING_Daily_{daily[0]['d']}_{daily[-1]['d']}.xlsx")
    try:
        wb.save(path)
        print(f"LONG STAYING daily: {path}")
    except PermissionError:
        print(f"!! ไฟล์เปิดอยู่ ไม่ได้บันทึก: {path}")


def archive_gateout(src):
    day = gate_day(src)
    folder = os.path.join(HERE, day)
    os.makedirs(folder, exist_ok=True)
    dest = os.path.join(folder, os.path.basename(src))
    if os.path.abspath(src) != os.path.abspath(dest):
        shutil.move(src, dest)
    rows = gateout_rows(dest)
    path = os.path.join(folder, f"GATE_OUT_Report_{day}.xlsx")
    counts = write_gateout_xlsx(rows, day, os.path.basename(dest), path)
    print(f"Gate out {day}: {len(rows)} ตู้ปล่อยออก (OEP/OEV/OEW)")
    print(f"  - {os.path.basename(dest)} (ไฟล์ต้นฉบับ)")
    print(f"  - GATE_OUT_Report_{day}.xlsx  " + " · ".join(f"{k}: {v:,}" for k, v in counts.items()))


# กลุ่มวันตามสี (เหมือนใน dashboard): ชื่อชีท, ช่วงวัน, สีพื้น, สีตัวเลข Days
DAY_BUCKETS = [
    ("0-30 วัน", 0, 30, "DCFCE7", "15803D"),
    ("31-99 วัน", 31, 99, "FFEDD5", "C2410C"),
    ("100-199 วัน", 100, 199, "FEE2E2", "B91C1C"),
    ("200+ วัน", 200, 10**9, "FFF59D", "B91C1C"),
]


def write_summary_xlsx(rows, meta, path):
    """Excel สรุป: ชีท 'สรุป' + แยกชีทละกลุ่มวัน (0-30 / 31-99 / 100-199 / 200+)"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="BFC5CE")
    border = Border(top=thin, left=thin, bottom=thin, right=thin)
    head_fill = PatternFill("solid", fgColor="1F5FBF")
    total_fill = PatternFill("solid", fgColor="EEF1F5")

    def style_head(cells):
        for c in cells:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center")
            c.border = border

    in_bucket = {b[0]: [r for r in rows if b[1] <= r["d"] <= b[2]] for b in DAY_BUCKETS}

    wb = Workbook()
    ws = wb.active
    ws.title = "สรุป"
    ws["A1"] = "EQ STAYING — สรุปตามกลุ่มวัน"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"ไฟล์: {meta['source']} · EQ Date ล่าสุด: {meta['lastEq']} · ทั้งหมด {len(rows):,} ตู้"
    ws["A2"].font = Font(color="6B7684")

    def table(start_row, title, keys, key_fn):
        ws.cell(start_row, 1, title).font = Font(bold=True, size=12)
        r0 = start_row + 1
        hdr = [""] + [b[0] for b in DAY_BUCKETS] + ["รวม"]
        for j, h in enumerate(hdr, 1):
            ws.cell(r0, j, h)
        style_head(ws[r0][:len(hdr)])
        for b, cell in zip(DAY_BUCKETS, ws[r0][1:len(DAY_BUCKETS) + 1]):
            cell.fill = PatternFill("solid", fgColor=b[3])
            cell.font = Font(bold=True, color=b[4])
        for i, k in enumerate(keys + ["รวม"], 1):
            rr = r0 + i
            vals = [sum(1 for x in in_bucket[b[0]] if k == "รวม" or key_fn(x) == k) for b in DAY_BUCKETS]
            for j, v in enumerate([k] + vals + [sum(vals)], 1):
                c = ws.cell(rr, j, v)
                c.border = border
                if k == "รวม" or j in (1, len(hdr)):
                    c.font = Font(bold=True)
                if k == "รวม":
                    c.fill = total_fill
        return r0 + len(keys) + 3

    nxt = table(4, "แยกตาม Area", ["BKK", "LCH"], lambda x: x["ar"])
    nxt = table(nxt, "แยกตาม Group", ["AV", "DMG", "OTHER"], lambda x: x["gp"])
    sizes = sorted({x["st"] for x in rows}, key=lambda s: -sum(1 for x in rows if x["st"] == s))
    nxt = table(nxt, "แยกตาม Size/Type", sizes, lambda x: x["st"])
    locs = sorted({(x["ar"], x["loc"]) for x in rows})
    table(nxt, "แยกตาม Location", [l for _, l in locs], lambda x: x["loc"])
    ws.column_dimensions["A"].width = 16
    for c in "BCDEF":
        ws.column_dimensions[c].width = 13

    head = ["#", "Area", "Location", "Days", "Container No", "Size/Type", "Group",
            "Move Code", "EQ Date", "Built Year", "RF Brand"]
    widths = [6, 8, 11, 8, 16, 10, 8, 11, 12, 11, 13]
    for name, lo, hi, fill_hex, font_hex in DAY_BUCKETS:
        sh = wb.create_sheet(name)
        sh.append(head)
        style_head(sh[1])
        fill = PatternFill("solid", fgColor=fill_hex)
        for i, r in enumerate(sorted(in_bucket[name], key=lambda x: -x["d"]), 1):
            sh.append([i, r["ar"], r["loc"], r["d"], r["cn"], r["st"], r["gp"],
                       r["mc"], r["eq"], r["by"] or "", r["rf"]])
            for c in sh[i + 1]:
                c.fill = fill
                c.border = border
            sh.cell(i + 1, 4).font = Font(bold=True, color=font_hex)
        for c, w in zip("ABCDEFGHIJK", widths):
            sh.column_dimensions[c].width = w
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = f"A1:K{len(in_bucket[name]) + 1}"
        sh.sheet_properties.tabColor = font_hex if name != "200+ วัน" else "FDE047"
    wb.save(path)
    return {b[0]: len(in_bucket[b[0]]) for b in DAY_BUCKETS}


def write_gateout_xlsx(rows, day, source, path):
    """รายงานตู้ปล่อยออก (OEP/OEV/OEW): ชีทสรุป + แยกชีทตามกลุ่มวันค้างก่อนปล่อย"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="BFC5CE")
    border = Border(top=thin, left=thin, bottom=thin, right=thin)

    def style_head(cells):
        for c in cells:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F5FBF")
            c.alignment = Alignment(horizontal="center")
            c.border = border

    known = [r for r in rows if r["d"] is not None]
    in_bucket = {b[0]: [r for r in known if b[1] <= r["d"] <= b[2]] for b in DAY_BUCKETS}
    unknown = [r for r in rows if r["d"] is None]

    wb = Workbook()
    ws = wb.active
    ws.title = "สรุป"
    ws["A1"] = f"รายงานตู้ปล่อยออกจากลาน (OEP / OEV / OEW) — {day}"
    ws["A1"].font = Font(bold=True, size=14)
    avg = round(sum(r["d"] for r in known) / len(known), 1) if known else 0
    long_n = sum(1 for r in known if r["d"] >= 100)
    ws["A2"] = (f"ไฟล์: {source} · ปล่อยออก {len(rows):,} ตู้ · วันค้างเฉลี่ย {avg} วัน · "
                f"ค้าง ≥100 วัน {long_n:,} ตู้" + (f" · ไม่พบวันเข้าลาน {len(unknown)} ตู้" if unknown else ""))
    ws["A2"].font = Font(color="6B7684")

    def table(start_row, title, keys, key_fn):
        ws.cell(start_row, 1, title).font = Font(bold=True, size=12)
        r0 = start_row + 1
        hdr = [""] + [b[0] for b in DAY_BUCKETS] + ["รวม"]
        for j, h in enumerate(hdr, 1):
            ws.cell(r0, j, h)
        style_head(ws[r0][:len(hdr)])
        for b, cell in zip(DAY_BUCKETS, ws[r0][1:len(DAY_BUCKETS) + 1]):
            cell.fill = PatternFill("solid", fgColor=b[3])
            cell.font = Font(bold=True, color=b[4])
        for i, k in enumerate(keys + ["รวม"], 1):
            rr = r0 + i
            vals = [sum(1 for x in in_bucket[b[0]] if k == "รวม" or key_fn(x) == k) for b in DAY_BUCKETS]
            for j, v in enumerate([k] + vals + [sum(vals)], 1):
                c = ws.cell(rr, j, v)
                c.border = border
                if k == "รวม" or j in (1, len(hdr)):
                    c.font = Font(bold=True)
                if k == "รวม":
                    c.fill = PatternFill("solid", fgColor="EEF1F5")
        return r0 + len(keys) + 3

    nxt = table(4, "แยกตาม Move (การปล่อย)", sorted({r["mv"] for r in known}), lambda x: x["mv"])
    nxt = table(nxt, "แยกตาม Group ก่อนปล่อย (AV / DMG)", ["AV", "DMG", "OTHER"], lambda x: x["pgp"])
    nxt = table(nxt, "แยกตาม Size/Type", sorted({r["st"] for r in known}), lambda x: x["st"])
    table(nxt, "แยกตาม Location", sorted({r["loc"] for r in known}), lambda x: x["loc"])
    ws.column_dimensions["A"].width = 18
    for c in "BCDEF":
        ws.column_dimensions[c].width = 13

    head = ["#", "Gate Date", "Time", "Move", "Area", "Location", "Days ค้าง", "Container No", "Size/Type",
            "EQ Date (เข้าลาน)", "Move ก่อนปล่อย", "Group", "Built Year", "POD", "Voyage"]
    widths = [6, 12, 7, 7, 7, 10, 9, 16, 10, 16, 14, 8, 10, 9, 9]
    sheets = [(b[0], in_bucket[b[0]], b[3], b[4]) for b in DAY_BUCKETS]
    if unknown:
        sheets.append(("ไม่พบวันเข้าลาน", unknown, "EEF1F5", "6B7684"))
    for name, lst, fill_hex, font_hex in sheets:
        sh = wb.create_sheet(name)
        sh.append(head)
        style_head(sh[1])
        fill = PatternFill("solid", fgColor=fill_hex)
        for i, r in enumerate(lst, 1):
            sh.append([i, r["gd"], r["gt"], r["mv"], r["ar"], r["loc"], r["d"] if r["d"] is not None else "",
                       r["cn"], r["st"], r["eq"], r["pmc"], r["pgp"], r["by"] or "", r["pod"], r["voy"]])
            for c in sh[i + 1]:
                c.fill = fill
                c.border = border
            sh.cell(i + 1, 7).font = Font(bold=True, color=font_hex)
        for col, w in zip("ABCDEFGHIJKLMNO", widths):
            sh.column_dimensions[col].width = w
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = f"A1:O{len(lst) + 1}"
        sh.sheet_properties.tabColor = "FDE047" if name == "200+ วัน" else font_hex
    wb.save(path)
    counts = {b[0]: len(in_bucket[b[0]]) for b in DAY_BUCKETS}
    if unknown:
        counts["ไม่พบวันเข้าลาน"] = len(unknown)
    return counts


def iso_week(d):
    """'2026-10-22' -> ('2026-W43', '19/10–25/10')"""
    from datetime import timedelta
    dt = datetime.strptime(d, "%Y-%m-%d")
    y, w, wd = dt.isocalendar()
    mon = dt - timedelta(days=wd - 1)
    sun = mon + timedelta(days=6)
    return f"{y}-W{w:02d}", f"{mon:%d/%m}–{sun:%d/%m}"


def write_gate_summaries(rows, meta, gate, ldaily=()):
    """ตัดยอดและสรุปตู้ปล่อยออกรายวัน
    - EQ STAYING/GATE OUT/GATE_OUT_Summary.xlsx : รายสัปดาห์ / รายวัน (พร้อมยอดคงเหลือ) / แต่ละลาน / รายการ
    - EQ STAYING/<วันที่>/GATE_OUT_<วันที่>.xlsx : สรุปและรายการตู้ปล่อยออกของวันนั้น"""
    if not gate:
        return
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="BFC5CE")
    border = Border(top=thin, left=thin, bottom=thin, right=thin)
    ld = {x["d"]: x for x in ldaily}
    days = sorted({g["gd"] for g in gate})
    days = sorted(set(days) | {h for h in YARD_HOLIDAYS if days[0] <= h <= days[-1]})
    yards = sorted({(g["loc"], g["yard"]) for g in gate})

    def bucket_counts(lst):
        k = [g for g in lst if g["d"] is not None]
        return [sum(1 for g in k if lo <= g["d"] <= hi) for _, lo, hi, _, _ in DAY_BUCKETS]

    def stats(lst):
        k = [g["d"] for g in lst if g["d"] is not None]
        return [len(lst)] + bucket_counts(lst) + [round(sum(k) / len(k), 1) if k else "", max(k) if k else ""]

    def table(ws, start, title, head, data, total=None):
        ws.cell(start, 1, title).font = Font(bold=True, size=12)
        r0 = start + 1
        for j, h in enumerate(head, 1):
            c = ws.cell(r0, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F5FBF")
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border = border
        for b, j in zip(DAY_BUCKETS, range(len(head))):
            pass
        for i, vals in enumerate(data + ([total] if total else []), 1):
            for j, v in enumerate(vals, 1):
                c = ws.cell(r0 + i, j, v)
                c.border = border
                if total and i == len(data) + 1:
                    c.font = Font(bold=True)
                    c.fill = PatternFill("solid", fgColor="EEF1F5")
        return r0 + len(data) + (1 if total else 0) + 2

    def list_sheet(wb, name, lst):
        sh = wb.create_sheet(name)
        head = ["No.", "Gate Date", "Time", "ลาน", "Location", "Area", "วันค้าง", "Container No", "Size/Type",
                "EQ Date (เข้าลาน)", "Move ก่อนปล่อย", "Group", "Type"]
        sh.append(head)
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0F766E")
            c.border = border
        fills = {b[0]: PatternFill("solid", fgColor=b[3]) for b in DAY_BUCKETS}
        for i, g in enumerate(sorted(lst, key=lambda x: (-(x["d"] if x["d"] is not None else -1), x["gd"])), 1):
            sh.append([i, g["gd"], g["gt"], g["yard"], g["loc"], g["ar"], g["d"] if g["d"] is not None else "",
                       g["cn"], g["st"], g["eq"], g["pmc"], g["pgp"], g["typ"]])
            b = next((b for b in DAY_BUCKETS if g["d"] is not None and b[1] <= g["d"] <= b[2]), None)
            for c in sh[i + 1]:
                c.border = border
                if b:
                    c.fill = fills[b[0]]
            if b:
                sh.cell(i + 1, 7).font = Font(bold=True, color=b[4])
        for col, w in zip("ABCDEFGHIJKLM", [6, 11, 7, 11, 9, 6, 8, 15, 9, 15, 13, 7, 10]):
            sh.column_dimensions[col].width = w
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = f"A1:M{len(lst) + 1}"

    bhead = [b[0] for b in DAY_BUCKETS]
    shead = ["ปล่อยออก"] + bhead + ["เฉลี่ย (วัน)", "Max"]

    # ---------- ภาพรวมทั้งช่วง ----------
    outdir = os.path.join(HERE, "GATE OUT")
    os.makedirs(outdir, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "สรุป"
    ws["A1"] = f"สรุปตู้ปล่อยออก (GATE OUT จากรีพอตลาน) {days[0]} – {days[-1]}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"ยอดในลาน/LONG STAYING รายวันจาก STAYING ใน STOCK DAILY (ฉบับล่าสุดที่ไม่เกินวันนั้น) ตัดตู้ที่ปล่อยออกตามรีพอตลาน · "
                f"สร้างเมื่อ {meta['generated']}")
    ws["A2"].font = Font(color="6B7684")
    weeks = {}
    for d in days:
        weeks.setdefault(iso_week(d), []).append(d)
    nxt = table(ws, 4, "รายสัปดาห์ (ISO จันทร์–อาทิตย์)", ["สัปดาห์", "ช่วงวันที่", "จำนวนวัน"] + shead,
                [[w, rng, len(ds)] + stats([g for g in gate if g["gd"] in ds]) for (w, rng), ds in weeks.items()],
                ["รวม", "", len(days)] + stats(gate))
    daily = []
    for d in days:
        x = ld.get(d, {})
        daily.append([d, iso_week(d)[0]] + stats([g for g in gate if g["gd"] == d]) +
                     [x.get("snap", ""), x.get("inYard", ""), x.get("start", ""), x.get("out", ""), x.get("end", "")])
    nxt = table(ws, nxt, "รายวัน + ตัดยอด (ยอดคงเหลือจากยอดตั้งต้น)",
                ["วันที่ปล่อย", "สัปดาห์"] + shead + ["STAYING ที่ใช้", "ตู้ในลานต้นวัน", "LONG ต้นวัน", "LONG ปล่อยวันนี้", "LONG คงเหลือ"], daily)
    table(ws, nxt, "แต่ละลาน", ["ลาน", "Location"] + shead,
          [[y, l] + stats([g for g in gate if g["loc"] == l]) for l, y in yards], ["รวม", ""] + stats(gate))
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 13
    for col in "CDEFGHIJKLMN":
        ws.column_dimensions[col].width = 11
    list_sheet(wb, "ตู้ปล่อยออกทั้งหมด", gate)
    list_sheet(wb, "ค้าง ≥100 วัน", [g for g in gate if g["d"] is not None and g["d"] >= 100])
    path = os.path.join(outdir, "GATE_OUT_Summary.xlsx")
    try:
        wb.save(path)
        print(f"Gate summary: {path}")
    except PermissionError:
        print(f"!! ไฟล์เปิดอยู่ ไม่ได้บันทึก: {path}")

    # ---------- รายวัน: <วันที่>/GATE_OUT_<วันที่>.xlsx ----------
    for row in daily:
        d = row[0]
        lst = [g for g in gate if g["gd"] == d]
        wb = Workbook()
        ws = wb.active
        ws.title = "สรุป"
        ws["A1"] = f"ตู้ปล่อยออกจากลาน วันที่ {d[8:10]}/{d[5:7]}/{d[:4]} ({iso_week(d)[0]})" + (f" — {YARD_HOLIDAYS[d]}" if d in YARD_HOLIDAYS else "")
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (f"ปล่อยออก {len(lst):,} ตู้ · LONG STAYING ต้นวัน {row[-3]} · ปล่อยวันนี้ {row[-2]} · คงเหลือ {row[-1]} "
                    f"(STAYING {row[-5]})")
        ws["A2"].font = Font(color="6B7684")
        table(ws, 4, "แต่ละลาน", ["ลาน", "Location"] + shead,
              [[y, l] + stats([g for g in lst if g["loc"] == l]) for l, y in yards if any(g["loc"] == l for g in lst)],
              ["รวม", ""] + stats(lst))
        ws.column_dimensions["A"].width = 14
        for col in "BCDEFGHI":
            ws.column_dimensions[col].width = 11
        list_sheet(wb, "ตู้ปล่อยออก", lst)
        folder = os.path.join(HERE, d)
        os.makedirs(folder, exist_ok=True)
        try:
            wb.save(os.path.join(folder, f"GATE_OUT_{d}.xlsx"))
        except PermissionError:
            print(f"!! ไฟล์เปิดอยู่ ไม่ได้บันทึก: GATE_OUT_{d}.xlsx")
    last = daily[-1]
    print(f"  ตัดยอดถึง {last[0]}: LONG ต้นวัน {last[-3]} · ปล่อย {last[-2]} · LONG คงเหลือ {last[-1]}"
          f" · ไฟล์รายวัน {len(daily)} วัน")


def archive(src, rows, meta):
    """เก็บไฟล์ของรอบนี้ไว้ในโฟลเดอร์วันที่ เช่น EQ STAYING/2026-10-05/
    (ถ้าไฟล์ต้นฉบับอยู่ในโฟลเดอร์วันที่อยู่แล้ว จะใช้โฟลเดอร์นั้น)"""
    parent = os.path.dirname(os.path.abspath(src))
    if os.path.dirname(parent) == HERE and re.fullmatch(r"\d{4}-\d{2}-\d{2}", os.path.basename(parent)):
        day = os.path.basename(parent)
    else:
        day = datetime.now().strftime("%Y-%m-%d")
    folder = os.path.join(HERE, day)
    os.makedirs(folder, exist_ok=True)
    # ไฟล์ต้นฉบับ: ย้ายเข้าโฟลเดอร์วัน (ถ้ายังไม่อยู่ในนั้น)
    dest_src = os.path.join(folder, os.path.basename(src))
    if os.path.abspath(src) != os.path.abspath(dest_src):
        shutil.move(src, dest_src)
    shutil.copy2(OUT, os.path.join(folder, f"EQ_STAYING_Dashboard_{day}.html"))
    counts = write_summary_xlsx(rows, meta, os.path.join(folder, f"EQ_STAYING_Summary_{day}.xlsx"))
    print(f"Archive: {folder}")
    print(f"  - {os.path.basename(src)} (ไฟล์ต้นฉบับ)")
    print(f"  - EQ_STAYING_Dashboard_{day}.html")
    print(f"  - EQ_STAYING_Summary_{day}.xlsx  " + " · ".join(f"{k}: {v:,}" for k, v in counts.items()))
    return dest_src


if __name__ == "__main__":
    main()
